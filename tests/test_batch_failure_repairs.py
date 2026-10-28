"""No paid calls: regressions for failures observed in the ten-resume benchmark."""
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import pytest
from docx import Document
from local_api.services.word_preparation import materialize_textboxes
from local_api.services.text_extraction_service import extract_text
from local_api.services.review_chunking import audit_decisions, audit_visible_sources, ChunkError
from local_api.services.record_subfields import responsibility_fragment


def textbox_docx():
    doc=Document();doc.add_paragraph('Experience');doc.add_paragraph('After box')
    data=BytesIO();doc.save(data);out=BytesIO()
    w='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    mc='{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    with ZipFile(data) as source,ZipFile(out,'w') as dest:
        for item in source.infolist():
            value=source.read(item.filename)
            if item.filename=='word/document.xml':
                root=ET.fromstring(value);anchor=root.find(w+'body').find(w+'p')
                alternate=ET.SubElement(anchor,mc+'AlternateContent')
                for tag in ['Choice','Fallback']:
                    branch=ET.SubElement(alternate,mc+tag)
                    box=ET.SubElement(branch,w+'txbxContent')
                    for text in ['Example Labs','Sep 21 to Current Role: Engineer']:
                        ET.SubElement(ET.SubElement(ET.SubElement(box,w+'p'),w+'r'),w+'t').text=text
                value=ET.tostring(root)
            dest.writestr(item,value)
    return out.getvalue()


def test_textbox_materialization_preserves_order_and_deduplicates_alternate_shape():
    original=textbox_docx();before=bytes(original)
    fixed,count=materialize_textboxes(original)
    assert count==1 and original==before
    paragraphs=[p.text for p in Document(BytesIO(fixed)).paragraphs]
    assert paragraphs==['Experience','Example Labs','Sep 21 to Current Role: Engineer','After box']
    assert materialize_textboxes(fixed)==(fixed,0)
    extracted=extract_text(original,'test.docx')
    assert extracted.text.count('Example Labs')==1
    assert any(e.metadata.get('materialized_word_textboxes')==1 for e in extracted.elements)


def audit_chunk(kind='legacy_value',value=None):
    return {'targets':[{'id':'target','record_key':'experience/job','kind':kind,'value':value}],
            'context_by_record':{'experience/job':[{'source_id':'own'}]}}


def test_audit_missing_legacy_value_may_report_absence_without_fabricated_citation():
    reply=json.dumps({'checked_ids':['target'],'findings':[{'field':'target','issue':'not established','source_ids':[]}]})
    assert len(audit_decisions(audit_chunk(),reply))==1
    with pytest.raises(ChunkError,match='foreign_audit_source'):
        audit_decisions(audit_chunk(value='populated'),reply)


def test_audit_foreign_quote_still_rejected_and_boundary_context_keeps_owners():
    chunk=audit_chunk('unresolved_field')
    reply=json.dumps({'checked_ids':['target'],'findings':[{'field':'target','issue':'wrong','source_ids':['foreign']}]})
    with pytest.raises(ChunkError,match='foreign_audit_source'):audit_decisions(chunk,reply)
    visible=audit_visible_sources({'targets':[], 'context_by_record':{'boundary':[
        {'source_id':'left','parent':'experience','record_id':'a'},
        {'source_id':'right','parent':'experience','record_id':'b'}]}})
    assert visible['experience/a']=={'left'} and visible['experience/b']=={'right'}


def test_capability_sentence_is_not_employer_evidence():
    from local_api.services.record_relationships import field_kinds
    text='Familiar with JSON based REST Web Services and Example Cloud Services.'
    assert responsibility_fragment(text)
    assert 'employer' not in field_kinds(text,'experience')


def test_empty_experience_record_is_retained_as_unmapped_source_not_fake_entity():
    from test_llm_schema_mapping import fixture
    from local_api.services.llm_schema_mapping import parse_review
    req,payload=fixture()
    record=payload['sections'][0]['records'][0]
    record['fields']={k:[] for k in record['fields']}
    record['missing']={k:'unmapped' for k in record['fields']}
    # Duration would recover a field from this header, so use a source with no dates.
    block=req.sections.sections[0].blocks[0]
    block.text='Shared responsibilities';block.end=len(block.text)
    record['source_ids']=[f'e0:0:{len(block.text)}']
    hierarchy,revised,_=parse_review(json.dumps(payload),req.sections)
    assert len(hierarchy[0].records)==1
    assert all(0 not in entry.block_indices for entry in revised.sections[0].entries)
    assert any(q.source_id==record['source_ids'][0] for q in hierarchy[0].residual_blocks)


def test_invalid_json_in_one_batch_does_not_discard_another_valid_batch():
    from test_llm_schema_mapping import fixture,configured
    from local_api.schemas.sections import ResumeSection,SectionBlock
    from local_api.schemas.extraction import ExtractedElement
    from local_api.services.llm_schema_mapping import map_with_field_review
    req,reply=fixture()
    text='Experienced software professional.'
    req.elements.append(ExtractedElement(type='NarrativeText',text=text,page_number=1,metadata={}))
    req.sections.sections.insert(0,ResumeSection(category='summary',label='Summary',text=text,blocks=[SectionBlock(
        source_element_index=2,start=0,end=len(text),text=text,page_number=1,assignment='semantic',lane='main',reason='fixture')]))
    def transport(settings,messages):
        categories=list(json.loads(messages[1]['content'])['registry'])
        return '{invalid JSON' if categories==['summary'] else json.dumps(reply)
    out=map_with_field_review(req,settings=configured(recovery_attempts=0),transport=transport)
    assert out.meta['field_review']['status']=='partial'
    assert out.meta['field_review']['fallback_categories']==['summary']
    assert [(r.company,r.role) for r in out.data.experience]==[('Acme Ltd','Engineer'),('Beta Ltd','Analyst')]
    assert out.meta['field_review']['diagnostics'][0]['recovered'] is False
    assert out.meta['mapping_result']['fallback_used'] is True


def test_terminal_json_repair_never_invents_values_or_relaxes_evidence():
    from local_api.services.llm_schema_mapping import repair_terminal_json_closer,map_with_field_review
    from test_llm_schema_mapping import fixture,configured
    req,payload=fixture()
    original=json.dumps(payload)
    cut=original.rfind('}',0,len(original)-1)
    damaged=original[:cut]+original[cut+1:]
    assert repair_terminal_json_closer(damaged)==payload
    out=map_with_field_review(req,settings=configured(recovery_attempts=0),transport=lambda *_:damaged)
    assert out.meta['field_review']['status']=='complete'
    assert out.meta['field_review']['diagnostics'][0]['recovered'] is True
    assert any(r['code']=='terminal_json_closer_inserted' for r in out.meta['field_review']['evidence_repairs'])
    payload['sections'][0]['records'][0]['fields']['company'][0]['value']='INVENTED EMPLOYER'
    bad=json.dumps(payload);cut=bad.rfind('}',0,len(bad)-1);bad=bad[:cut]+bad[cut+1:]
    rejected=map_with_field_review(req,settings=configured(recovery_attempts=0),transport=lambda *_:bad)
    assert rejected.meta['field_review']['status']=='failed'


@pytest.mark.parametrize('content',['{"name":"unfinished','{"name":}', '{"a" "b"}', '{"a": 1, "b":2'])
def test_ambiguous_or_truncated_json_is_not_reconstructed(content):
    from local_api.services.llm_schema_mapping import repair_terminal_json_closer
    assert repair_terminal_json_closer(content) is None


def test_word_table_paragraph_cuts_do_not_rewrite_text_or_guess_alignment():
    from local_api.services.word_preparation import annotate_table_paragraphs
    doc=Document();cell=doc.add_table(rows=1,cols=1).cell(0,0)
    cell.text='08/2024 – Present'
    cell.add_paragraph('Company - Example Labs')
    cell.add_paragraph('Engineer')
    source=BytesIO();doc.save(source)
    text='08/2024 – Present Company - Example Labs Engineer'
    elements=[{'type':'Table','text':text,'metadata':{}}]
    annotate_table_paragraphs(source.getvalue(),elements)
    assert elements[0]['text']==text
    spans=elements[0]['metadata']['word_source_paragraphs']
    assert [text[s['start']:s['end']] for s in spans]==['08/2024 – Present','Company - Example Labs','Engineer']
    unmatched=[{'type':'Table','text':text+' unsupported extra','metadata':{}}]
    annotate_table_paragraphs(source.getvalue(),unmatched)
    assert unmatched[0]['metadata']=={}


def test_technical_field_marker_and_prose_cannot_replace_established_company():
    from local_api.services.record_relationships import field_kinds,annotate_relationships
    from local_api.schemas.sections import ResumeSection,SectionBlock,SectionEntry
    assert 'employer' not in field_kinds('Technology -','experience')
    assert 'employer' not in field_kinds('Company -','experience')
    lines=['Example Pvt Ltd','Senior Engineer','Streamlined database performance and provided solutions for better operational efficiency.']
    s=ResumeSection(category='experience',label='Experience',blocks=[SectionBlock(source_element_index=i,
        start=0,end=len(t),text=t,page_number=1,lane='main',assignment='semantic',reason='fixture') for i,t in enumerate(lines)],
        entries=[SectionEntry(entry_id='job',block_indices=[0,1,2],lane='main',boundary_reason='fixture')])
    annotate_relationships(s)
    assert s.entries[0].fields['employer']==[0]
    assert s.entries[0].fields['responsibilities']==[2]


def test_explicit_company_client_header_does_not_use_client_as_employer():
    from local_api.services.record_subfields import employment_header_fields,labelled_fields
    from types import SimpleNamespace
    assert employment_header_fields('Example Labs, Client - Other Ltd')=={'company':'Example Labs'}
    fields=labelled_fields('experience',[SimpleNamespace(text='Company- Example Labs, Client - Other Ltd',assignment='semantic')])
    assert fields['company']==['Example Labs']


def test_actual_repaired_word_records_keep_two_jobs_and_four_project_names():
    from local_api.schemas.mapping import SchemaMappingRequest
    from local_api.services.resume_sections import build_sections
    from local_api.services.schema_mapping import map_schema
    root=Path('reports/resume-batch-repairs-20261027/table-verified')
    rows=json.loads((root/'results.json').read_text())
    for row in rows[3:5]:
        artifact=json.loads((root/row['artifact']).read_text())
        req=SchemaMappingRequest(elements=artifact['extraction']['elements'],sections=artifact['sections'])
        req.sections.sections=build_sections({s.category:s.blocks for s in req.sections.sections},req.elements)
        out=map_schema(req)
        assert [(r.company,r.role) for r in out.data.experience]==[('Velocity IT','DataScientist'),('Netafim','Software engineer')]
        assert [r.name for r in out.data.projects]==['Pre-screen Calling Agent','Question Identification System',
            'Clinical Trial Adverse Event Classifier','Automated Patient Condition Extraction from Clinical Conversations']


@pytest.mark.parametrize('text,expected',[
    ('Worked as a Technical Lead for Example Labs since June-2025 to Sep 2025.', {'role':'Technical Lead','company':'Example Labs'}),
    ('Example Labs : Dec 2023- till now', {'company':'Example Labs'}),
    ('Worked at Example Labs from JUN 2023 to JAN 2026', {'company':'Example Labs'}),
])
def test_explicit_alternate_employment_syntax(text,expected):
    from local_api.services.record_subfields import employment_header_fields
    from local_api.services.record_relationships import field_kinds
    assert employment_header_fields(text)==expected
    assert 'employer' in field_kinds(text,'experience')


def test_compound_heading_uses_schema_families_not_resume_specific_phrase():
    from local_api.services.resume_sections import ALIASES
    from local_api.services.resume_structure import heading_boundary
    assert heading_boundary('Projects and POCs',ALIASES)==('projects',None)
    assert heading_boundary('Projects and Education',ALIASES)==(None,None)
    category,cut=heading_boundary('Project 2: Search Portal',ALIASES)
    assert category=='projects' and 'Project 2: Search Portal'[cut:].strip()=='Search Portal'


def test_transport_failure_retains_preceding_valid_records():
    from test_llm_schema_mapping import fixture,configured
    from local_api.schemas.sections import ResumeSection,SectionBlock
    from local_api.schemas.extraction import ExtractedElement
    from local_api.services.llm_schema_mapping import map_with_field_review
    from local_api.services.llm_section_review import ReviewError
    req,payload=fixture();text='Professional summary.'
    req.elements.append(ExtractedElement(type='NarrativeText',text=text,page_number=1,metadata={}))
    req.sections.sections.append(ResumeSection(category='summary',label='Summary',text=text,blocks=[SectionBlock(
        source_element_index=2,start=0,end=len(text),text=text,page_number=1,assignment='semantic',lane='main',reason='fixture')]))
    calls=[]
    def transport(settings,messages):
        calls.append(1)
        if len(calls)>1:raise ReviewError('provider_unavailable')
        return json.dumps(payload)
    out=map_with_field_review(req,settings=configured(),transport=transport)
    assert len(calls)==2
    assert out.meta['field_review']['status']=='partial'
    assert [r.company for r in out.data.experience]==['Acme Ltd','Beta Ltd']
    assert out.meta['field_review']['diagnostics'][-1]['recovered'] is False


def test_actual_flattened_table_recovers_seven_distinct_employer_records():
    from local_api.schemas.mapping import SchemaMappingRequest
    from local_api.services.resume_sections import build_sections
    from local_api.services.schema_mapping import map_schema
    root=Path('reports/resume-batch-repairs-20261027/table-verified')
    row=json.loads((root/'results.json').read_text())[6]
    artifact=json.loads((root/row['artifact']).read_text())
    req=SchemaMappingRequest(elements=artifact['extraction']['elements'],sections=artifact['sections'])
    req.sections.sections=build_sections({s.category:s.blocks for s in req.sections.sections},req.elements)
    out=map_schema(req)
    assert len(out.data.experience)==7
    assert all(r.company for r in out.data.experience)
    assert 'Mantra' in out.data.experience[0].company
    assert [out.data.experience[i].company for i in [1,2,4]]==['NISG DELHI']*3
    assert 'Client' not in out.data.experience[-1].company


def test_established_experience_prefers_complete_dated_identity_over_duty_mentions():
    from local_api.services.record_subfields import refine_record
    from types import SimpleNamespace
    blocks=[SimpleNamespace(text='Example Labs (June 2023 – Jan 2026)',assignment='semantic',source_element_index=0,start=0,end=42),
            SimpleNamespace(text='for orchestration and cloud translation services.',assignment='semantic',source_element_index=1,start=0,end=48)]
    values,_=refine_record('experience',{'company':'','role':'','start_date':None,'end_date':None,'duration':None},blocks,header_fields={'company'})
    assert values['company']=='Example Labs'
    assert values['start_date']=='June 2023' and values['end_date']=='Jan 2026'


def test_project_subfield_titles_do_not_create_new_project_records():
    from local_api.services.resume_structure import entry_groups
    from local_api.schemas.sections import SectionBlock
    from local_api.schemas.extraction import ExtractedElement
    texts=['Project: 1','Client Name: Example','Roles & Responsibilities','Duration: Jun-22 to Sep-22 Roles & Responsibilities',
           'Built a platform.','Project: 2','Role: Engineer']
    elements=[ExtractedElement(type='Title',text=t,page_number=1,metadata={}) for t in texts]
    blocks=[SectionBlock(source_element_index=i,start=0,end=len(t),text=t,page_number=1,lane='main',assignment='heading_context',reason='fixture') for i,t in enumerate(texts)]
    groups=entry_groups(blocks,'projects',elements)
    assert [g['block_indices'] for g in groups]==[[0,1,2,3,4],[5,6]]


def test_client_role_header_does_not_create_an_employer():
    from local_api.services.record_subfields import employment_header_fields
    from local_api.services.record_relationships import field_kinds
    text='Client Example Bank – Java Developer'
    assert employment_header_fields(text)=={}
    assert 'employer' not in field_kinds(text,'experience')
