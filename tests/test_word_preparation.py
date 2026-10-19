from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile
import subprocess
import xml.etree.ElementTree as ET

import pytest
from docx import Document
from local_api.services.word_preparation import convert_legacy_word, repair_null_image_relationships
from local_api.services.text_extraction_service import extract_text
from local_api.utils.file_validation import ExtractionError


def broken_image_docx(kind='image',target='../NULL'):
    document=Document();document.add_paragraph('Work experience');document.add_paragraph('Engineer at Example Systems')
    stream=BytesIO();document.save(stream);out=BytesIO()
    with ZipFile(stream) as source,ZipFile(out,'w') as dest:
        for item in source.infolist():
            value=source.read(item.filename)
            if item.filename=='word/_rels/document.xml.rels':
                root=ET.fromstring(value)
                ET.SubElement(root,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',{
                    'Id':'rIdBroken','Type':'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'+kind,'Target':target})
                value=ET.tostring(root)
            dest.writestr(item,value)
    return out.getvalue()


def test_null_image_recovery_preserves_document_xml_and_extracts_text():
    original=broken_image_docx();repaired,count=repair_null_image_relationships(original)
    assert count==1
    with ZipFile(BytesIO(original)) as a,ZipFile(BytesIO(repaired)) as b:
        assert a.read('word/document.xml')==b.read('word/document.xml')
    result=extract_text(original,'resume.docx')
    assert 'Engineer at Example Systems' in result.text
    assert all(e.metadata['removed_null_image_relationships']==1 for e in result.elements)


@pytest.mark.parametrize(('kind','target'),[('header','../NULL'),('image','media/missing.png')])
def test_non_null_images_and_text_relationships_are_not_removed(kind,target):
    original=broken_image_docx(kind,target)
    assert repair_null_image_relationships(original)==(original,0)


def test_legacy_conversion_uses_private_profile(monkeypatch):
    profiles=[]
    def run(command,**kwargs):
        profiles.append(command[1]);assert command[1].startswith('-env:UserInstallation=file:///')
        source=Path(command[-1]);assert source.read_bytes()==b'legacy';source.with_suffix('.docx').write_bytes(b'converted')
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(subprocess,'run',run)
    assert convert_legacy_word(b'legacy')==b'converted'
    assert convert_legacy_word(b'legacy')==b'converted'
    assert profiles[0]!=profiles[1]


def test_failed_legacy_conversion_has_clear_document_error(monkeypatch):
    monkeypatch.setattr(subprocess,'run',lambda *a,**k:SimpleNamespace(returncode=77))
    with pytest.raises(ExtractionError,match='could not be converted') as exc:convert_legacy_word(b'bad')
    assert exc.value.status==422
