"""New extraction contract: real formats, with mocks restricted to failure/unit tests."""
from io import BytesIO
from pathlib import Path
import shutil

import pytest
from fastapi.testclient import TestClient
from starlette.datastructures import UploadFile
from local_api.app import app
from local_api.services import text_extraction_service as service
from local_api.utils.file_validation import ExtractionError

client = TestClient(app)
FIXTURES = Path(__file__).parent / 'fixtures'
URL = '/api/v1/extract/text'


def upload(name, payload, **data):
    return client.post(URL, files={'file': (name, payload)}, data=data)


def test_txt_schema_and_safe_name():
    response = upload('../../note.txt', b'This is a text document. It contains a useful sentence.')
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['success'] is True
    assert result['filename'] == 'note.txt'
    assert result['file_type'] == 'text/plain'
    assert 'useful sentence' in result['text']
    assert result['metadata']['total_elements'] == len(result['elements'])
    assert result['metadata']['page_count'] is None
    assert all(e['metadata']['filename'] == 'note.txt' for e in result['elements'])
    assert all('file_directory' not in e['metadata'] for e in result['elements'])


@pytest.mark.parametrize('strategy', ['fast', 'auto'])
def test_pdf_real(strategy):
    response = upload('digital.pdf', (FIXTURES / 'digital-multipage.pdf').read_bytes(), strategy=strategy)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['text'].strip()
    assert result['metadata']['page_numbers']
    assert result['file_type'] == 'application/pdf'


def test_docx_real():
    from docx import Document
    document = Document()
    document.add_heading('Integration document', 0)
    document.add_paragraph('This is a real Word document with extractable content.')
    stream = BytesIO()
    document.save(stream)
    response = upload('sample.docx', stream.getvalue())
    assert response.status_code == 200, response.text
    assert 'Integration document' in response.json()['text']


@pytest.mark.parametrize('filename,payload', [
    ('sample.md', b'# Heading\n\nMarkdown content for local extraction.'),
    ('sample.html', b'<html><body><h1>Heading</h1><p>HTML content for local extraction.</p></body></html>'),
    ('sample.csv', b'name,value\napple,10\norange,20\n'),
])
def test_textual_formats_real(filename, payload):
    response = upload(filename, payload)
    assert response.status_code == 200, response.text
    assert response.json()['text'].strip()


def test_xlsx_real():
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(['Name', 'Value'])
    workbook.active.append(['Example', 10])
    stream = BytesIO()
    workbook.save(stream)
    response = upload('sample.xlsx', stream.getvalue())
    assert response.status_code == 200, response.text
    assert 'Example' in response.json()['text']


def test_pptx_real():
    from pptx import Presentation
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = 'Presentation sample'
    stream = BytesIO()
    presentation.save(stream)
    response = upload('sample.pptx', stream.getvalue())
    assert response.status_code == 200, response.text
    assert 'Presentation sample' in response.json()['text']


def test_omit_elements():
    response = upload('note.txt', b'Text without structured elements.', include_elements='false')
    assert response.status_code == 200, response.text
    assert response.json()['elements'] is None
    assert response.json()['metadata']['total_elements'] > 0


@pytest.mark.parametrize('name,payload,status', [
    ('empty.txt', b'', 422), ('program.exe', b'MZ\x00\x01', 415),
    ('fake.pdf', b'\x89PNG\r\n\x1a\ninvalid', 422),
    ('broken.pdf', b'%PDF-1.4\ninvalid PDF content', 422),
    ('renamed.txt', b'\x7fELF\x02\x01\x01\x00' + b'\x00' * 300, 415),
])
def test_invalid_documents(name, payload, status):
    response = upload(name, payload, strategy='fast' if name.endswith('.pdf') else 'auto')
    assert response.status_code == status, response.text


def test_strategy_validation():
    assert upload('note.txt', b'Text', strategy='hi_res').status_code == 422
    assert upload('note.txt', b'Text', strategy='invalid').status_code == 422


def test_upload_limit(monkeypatch):
    monkeypatch.setenv('MAX_UPLOAD_BYTES', '8')
    assert upload('note.txt', b'123456789').status_code == 413


def test_missing_file():
    assert client.post(URL).status_code == 400


@pytest.mark.parametrize('error,status', [(ImportError('secret local path'), 503),
                                         (ValueError('secret content'), 422),
                                         (RuntimeError('secret content'), 500)])
def test_extraction_errors_unit(monkeypatch, error, status, caplog):
    def fail(**kwargs):
        raise error
    monkeypatch.setattr(service, 'partition', fail)
    response = upload('note.txt', b'This is plain text.')
    assert response.status_code == status
    assert 'secret' not in response.text
    assert 'secret' not in caplog.text


@pytest.mark.parametrize('failure', [False, True])
def test_upload_cleanup(monkeypatch, failure):
    closed = []
    original = UploadFile.close
    async def close(self):
        await original(self)
        closed.append(self.file.closed)
    monkeypatch.setattr(UploadFile, 'close', close)
    if failure:
        monkeypatch.setenv('MAX_UPLOAD_BYTES', '1')
    response = upload('note.txt', b'Temporary upload cleanup test.')
    assert response.status_code == (413 if failure else 200)
    assert closed and all(closed)


@pytest.mark.skipif(not shutil.which('tesseract') or not shutil.which('pdftoppm'), reason='System OCR dependencies unavailable')
def test_scanned_pdf_ocr_real():
    response = upload('scanned.pdf', (FIXTURES / 'scanned.pdf').read_bytes(), strategy='ocr_only')
    assert response.status_code == 200, response.text
    assert response.json()['text'].strip()


@pytest.mark.skipif(not shutil.which('tesseract'), reason='Tesseract unavailable')
def test_image_ocr_real():
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new('RGB', (1000, 200), 'white')
    font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 36)
    ImageDraw.Draw(image).text((30, 60), 'Local image extraction example', fill='black', font=font)
    stream = BytesIO()
    image.save(stream, format='PNG')
    response = upload('sample.png', stream.getvalue(), strategy='ocr_only')
    assert response.status_code == 200, response.text
    assert 'extraction' in response.json()['text'].lower()


def test_direct_service_real():
    result = service.extract_text(b'This is directly callable extraction.', 'direct.txt')
    assert 'directly callable' in result.text
    assert result.filename == 'direct.txt'


def test_order_and_metadata_unit(monkeypatch):
    class Element:
        def __init__(self, text, page):
            self.text, self.page = text, page
        def to_dict(self):
            return {'type': 'NarrativeText', 'text': self.text,
                    'metadata': {'page_number': self.page, 'file_directory': '/private/path'}}
    monkeypatch.setattr(service, 'partition', lambda **kwargs: [Element('first', 3), Element('second', 1)])
    result = service.extract_text(b'Example text.', 'note.txt')
    assert result.text == 'first\n\nsecond'
    assert result.metadata.page_numbers == [1, 3]
    assert result.metadata.page_count is None
    assert all('file_directory' not in e.metadata for e in result.elements)


def test_missing_ocr_dependency_unit(monkeypatch):
    monkeypatch.setattr(service.shutil, 'which', lambda name: None)
    response = upload('scan.pdf', (FIXTURES / 'scanned.pdf').read_bytes(), strategy='ocr_only')
    assert response.status_code == 503
    assert 'Tesseract' in response.json()['detail']


def test_spooled_disk_cleanup_unit(monkeypatch):
    observed = []
    original = UploadFile.close
    async def close(self):
        rolled = getattr(self.file, '_rolled', False)
        await original(self)
        observed.append((rolled, self.file.closed))
    monkeypatch.setattr(UploadFile, 'close', close)
    monkeypatch.setattr(service, 'partition', lambda **kwargs: [])
    response = upload('large.txt', b'Example text. ' * 100_000)
    assert response.status_code == 200, response.text
    assert observed and all(rolled and closed for rolled, closed in observed)


def test_valid_blank_pdf_real():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    stream = BytesIO()
    writer.write(stream)
    response = upload('blank.pdf', stream.getvalue(), strategy='fast')
    assert response.status_code == 200, response.text
    assert response.json()['text'] == ''
    assert response.json()['metadata']['total_elements'] == 0
