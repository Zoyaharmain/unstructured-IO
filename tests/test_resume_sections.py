"""Real local BERT/K-means tests plus isolated boundary/error tests."""
from pathlib import Path
import json
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from local_api.app import app
from local_api.schemas.extraction import ExtractedElement
from local_api.schemas.sections import SectionRequest
from local_api.services import resume_sections as service
from local_api.utils.file_validation import ExtractionError

client = TestClient(app)


def request(texts, types=None, coordinates=None):
    return SectionRequest(extraction_verified=True, elements=[
        ExtractedElement(type=types[i] if types else 'NarrativeText', text=text, page_number=1,
                         metadata={'coordinates': coordinates[i]} if coordinates else {}) for i, text in enumerate(texts)])


@pytest.fixture(scope='module')
def engine():
    return service.get_engine()


def assert_preserved(original, response):
    blocks = [block for section in response.sections for block in section.blocks] + response.needs_review
    for index, element in enumerate(original.elements):
        parts = sorted([block for block in blocks if block.source_element_index == index], key=lambda block: block.start)
        cursor = 0
        for part in parts:
            assert part.start == cursor
            assert part.text == element.text[part.start:part.end]
            cursor = part.end
        assert cursor == len(element.text)
        assert ''.join(part.text for part in parts) == element.text
    assert response.metadata['all_source_text_preserved'] is True


def test_real_model_is_bert_and_actual_kmeans(engine):
    from sklearn.cluster import KMeans
    assert engine.encoder[0].auto_model.config.model_type == 'bert'
    assert isinstance(engine.headings.model, KMeans)
    assert isinstance(engine.content.model, KMeans)


def test_real_standard_resume(engine):
    texts = ['Jordan Taylor', 'jordan@example.com +1 555 0123 London', 'Professional summary',
             'Experienced software engineer with five years building reliable services.',
             'Work experience', 'Software Engineer at Acme Corporation, January 2020 to December 2023.',
             'Built Python APIs and reduced production latency by 35%.', 'Projects',
             'Expense Tracker: Built a personal application with React.', 'Education',
             'Bachelor of Science in Computer Science, Example University, GPA 3.8.', 'Skills',
             'Python, Java, SQL, Docker, AWS']
    req = request(texts)
    result = engine.detect(req)
    groups = {section.category: '\n'.join(block.text for block in section.blocks) for section in result.sections}
    assert 'Jordan Taylor' in groups['personal_information']
    assert texts[3] in groups['summary']
    assert texts[5] in groups['experience'] and texts[6] in groups['experience']
    assert texts[8] in groups['projects'] and texts[8] not in groups['experience']
    assert texts[10] in groups['education'] and texts[12] in groups['skills']
    assert_preserved(req, result)


def test_real_synonym_headings(engine):
    req = request(['Career overview', 'Results-oriented professional with strong leadership skills.',
                   'Employment history', 'Senior Developer, Example Ltd, remote position, 2021 to present.',
                   'Academic background', 'Master of Business Administration from a university, graduation year 2022.',
                   'Core competencies', 'Python, Java, SQL and Docker.'])
    result = engine.detect(req)
    groups = {section.category: [block.text for block in section.blocks] for section in result.sections}
    for index, category in [(0,'summary'),(2,'experience'),(4,'education'),(6,'skills')]:
        assert req.elements[index].text in groups[category]
        assert req.elements[index+1].text in groups[category]
    assert_preserved(req, result)


def coords(left, right):
    return {'layout_width': 1000, 'layout_height': 1400, 'points': [[left,100],[left,130],[right,130],[right,100]], 'system':'PixelSpace'}


@pytest.mark.parametrize('main_left', [270, 550])
def test_real_two_column_interleaved_resume(engine, main_left):
    texts = ['Skills', 'Work experience', 'Python Java SQL',
             'Software Engineer at Acme Corporation, January 2020 to December 2023.',
             'Docker AWS Git', 'Delivered production services for customers.',
             'Education', 'Projects', 'Bachelor of Science, Example University.',
             'Personal project: built an expense tracker with a GitHub repository.']
    coordinates = [coords(50,230) if i % 2 == 0 else coords(main_left,950) for i in range(len(texts))]
    req = request(texts, coordinates=coordinates)
    result = engine.detect(req)
    groups = {section.category: '\n'.join(block.text for block in section.blocks) for section in result.sections}
    assert texts[2] in groups['skills'] and texts[4] in groups['skills']
    assert texts[3] in groups['experience'] and texts[5] in groups['experience']
    assert texts[2] not in groups['experience'] and texts[4] not in groups['experience']
    assert texts[8] in groups['education'] and texts[9] in groups['projects']
    assert any('layout lanes' in warning for warning in result.warnings)
    assert_preserved(req, result)


def test_real_newlines_and_long_text_not_truncated(engine):
    text = 'Work experience\n' + ('Built production services for customers. ' * 130) + '\nProjects\nPersonal project: created an expense tracker.\n'
    req = request([text])
    result = engine.detect(req)
    assert result.metadata['total_units'] > 4
    assert_preserved(req, result)
    assert any(section.category == 'projects' for section in result.sections)


def test_real_ambiguous_input_is_not_forced(engine):
    req = request(['xyzzy qqqv zzzz 991827', '☃ ☂ ♞'])
    result = engine.detect(req)
    assert result.needs_review
    assert result.status == 'needs_review'
    assert_preserved(req, result)


def test_uncertain_heading_stops_context_unit(monkeypatch, engine):
    import numpy as np
    req = request(['Skills', 'Ambiguous heading', 'uncertain content'], types=['Title','Title','NarrativeText'])
    monkeypatch.setattr(engine.encoder, 'encode', lambda *args, **kwargs: np.zeros((3,384)))
    monkeypatch.setattr(engine.headings, 'score', lambda *args, **kwargs: [('skills',0,1,.5), ('projects',1,.7,.01), ('experience',2,.3,.01)])
    monkeypatch.setattr(engine.content, 'score', lambda *args, **kwargs: [('skills',0,.8,.4), ('projects',1,.3,.01), ('experience',2,.3,.01)])
    result = engine.detect(req)
    assert [block.text for block in result.sections[0].blocks] == ['Skills']
    assert len(result.needs_review) == 2
    assert_preserved(req, result)


def test_api_groups_before_optional_manual_verification():
    payload = {'elements':[{'type':'Title','text':'Skills','metadata':{}}]}
    response = client.post('/api/v1/resume/sections',json=payload)
    assert response.status_code == 200
    assert response.json()['extraction_verified'] is False
    payload['extraction_verified'] = False
    assert client.post('/api/v1/resume/sections',json=payload).status_code == 200
    payload['extraction_verified'] = True
    payload['elements'][0]['text'] = ' '
    assert client.post('/api/v1/resume/sections',json=payload).status_code == 422
    payload['elements'][0]['text'] = 'x' * 100001
    assert client.post('/api/v1/resume/sections',json=payload).status_code == 422


def test_missing_local_model_is_clear_unit(monkeypatch):
    monkeypatch.setattr(service, '_engine', None)
    def unavailable(*args, **kwargs):
        raise OSError('/private/model/path')
    monkeypatch.setattr(service, 'SectionEngine', unavailable)
    response = client.post('/api/v1/resume/sections', json=request(['Skills']).model_dump())
    assert response.status_code == 503
    assert 'warm_models.py --sections' in response.json()['detail']
    assert '/private/' not in response.text


def test_real_docx_extraction_to_sections(engine):
    from docx import Document
    document = Document()
    for title, text in [('Professional summary', 'Experienced software engineer with five years building services.'),
                        ('Work experience', 'Software Engineer at Acme Corporation, 2020 to 2023.'),
                        ('Projects', 'Personal project: created an expense tracker application.'),
                        ('Education', 'Bachelor of Science, Example University, 2020.'),
                        ('Technical skills', 'Python, Java, SQL and Docker.')]:
        document.add_heading(title, 1)
        document.add_paragraph(text)
    stream=BytesIO(); document.save(stream)
    extracted=client.post('/api/v1/extract/text', files={'file':('resume.docx',stream.getvalue())})
    assert extracted.status_code == 200, extracted.text
    payload = {'elements':extracted.json()['elements'], 'extraction_verified':True}
    response=client.post('/api/v1/resume/sections', json=payload)
    assert response.status_code == 200, response.text
    assert {'summary','experience','projects','education','skills'} <= {section['category'] for section in response.json()['sections']}
    from local_api.schemas.sections import SectionResponse
    assert_preserved(SectionRequest.model_validate(payload), SectionResponse.model_validate(response.json()))


def test_real_combined_headings_are_not_forced(engine):
    req = request(['Education and Certifications', 'Experience and Projects'], types=['Title', 'Title'])
    result = engine.detect(req)
    assert not result.sections
    assert len(result.needs_review) == 2
    assert set(result.needs_review[0].alternatives) == {'education', 'certifications'}
    assert set(result.needs_review[1].alternatives) == {'experience', 'projects'}
    assert_preserved(req, result)


def test_real_footer_is_not_mixed_into_experience(engine):
    req = request(['Work experience', 'Software Engineer at Acme Corporation, 2020 to 2023.',
                   'Page 1 of 2', '', 'Delivered production services for customers.'],
                  types=['Title','NarrativeText','Footer','PageBreak','NarrativeText'])
    result = engine.detect(req)
    experience = next(section for section in result.sections if section.category == 'experience')
    assert 'Page 1 of 2' not in [block.text for block in experience.blocks]
    assert req.elements[-1].text in [block.text for block in experience.blocks]
    assert {'Page 1 of 2', ''} <= {block.text for block in result.needs_review}
    assert_preserved(req, result)


def test_structure_inline_headings_and_project_fields(engine):
    req = request(['S K I L L S: Python Python SQL\nProjects: Expense tracker\n',
                   'Project name: Ledger', 'Technologies: Python, SQL',
                   'Achievement: Won an internal prize', 'Education: BSc, Example University'])
    result = engine.detect(req)
    groups = {section.category: ''.join(block.text for block in section.blocks) for section in result.sections}
    assert 'Python Python SQL' in groups['skills']
    assert 'Technologies: Python, SQL' in groups['projects']
    assert 'Achievement: Won an internal prize' in groups['projects']
    assert 'BSc, Example University' in groups['education']
    assert_preserved(req, result)


def test_structure_entries_do_not_cross_sections_or_lanes(engine):
    req = request(['Work experience', 'Engineer, Alpha Ltd, 2020 to 2022',
                   'Built Python services.', 'Engineer, Beta Ltd, 2022 to present',
                   'Delivered SQL services.', 'Projects', 'Project name: Ledger',
                   'Technologies: Python SQL', 'Project name: Portal', 'Technologies: React'],
                  types=['Title', 'Title', 'NarrativeText', 'Title', 'NarrativeText',
                         'Title', 'NarrativeText', 'NarrativeText', 'NarrativeText', 'NarrativeText'])
    result = engine.detect(req)
    sections = {section.category: section for section in result.sections}
    for category, expected in [('experience', 2), ('projects', 2)]:
        section = sections[category]
        assert len(section.entries) == expected
        indices = [index for entry in section.entries for index in entry.block_indices]
        assert len(indices) == len(set(indices))
        assert sorted(indices) == [i for i, block in enumerate(section.blocks) if block.assignment not in {'section_heading', 'structural_heading'}]
    assert_preserved(req, result)


def test_structure_does_not_promote_body_keywords(engine):
    req = request(['Projects', 'Project name: Contact Center',
                   'The contact center infrastructure uses Python.',
                   'Tools: Docker SQL', 'Skills: Java Java', '', 'Summary: AB'])
    result = engine.detect(req)
    projects = next(section for section in result.sections if section.category == 'projects')
    assert 'Tools: Docker SQL' in [block.text for block in projects.blocks]
    assert 'The contact center infrastructure uses Python.' in [block.text for block in projects.blocks]
    assert_preserved(req, result)


def test_merged_pdf_heading_prefixes(engine):
    req = request(['Summary Backend Developer building reliable APIs.',
                   'Experience Engineer at Acme (2020 - Present) Built Python APIs.',
                   'Projects Credit scoring application using Python.',
                   'Education B.Tech in Computer Science, Example University.',
                   'Certifications AWS Certified Developer.',
                   'Skills Python Python SQL'])
    result = engine.detect(req)
    groups = {s.category: ''.join(b.text for b in s.blocks) for s in result.sections}
    for category, element in zip(['summary','experience','projects','education','certifications','skills'], req.elements):
        assert groups[category] == element.text
    assert_preserved(req, result)


def test_leading_body_phrase_does_not_split_as_heading():
    from local_api.services.resume_structure import heading_boundary
    for text in ['Experience building reliable APIs.', 'Skills include Python and SQL.',
                 'The contact center project', 'Project name: Skills Portal']:
        assert heading_boundary(text, service.ALIASES) == (None, None)


def test_readable_work_history_has_employers_roles_and_wrapped_responsibilities(engine):
    req = request(['Profile Snapshot', 'Experienced finance analyst using Python and SQL.',
                   'Core Competencies', 'Accounting and reconciliation', 'Python SQL Python',
                   'Work History', 'Example India Private Limited | Hyderabad | Dec 2022 – Present',
                   'Senior Analyst – Fund Operations | Jan 2025 – Present',
                   'Managed reconciliations across multiple funds,', 'resolving discrepancies within SLA.',
                   'Built ETL workflows in Python.', 'Analyst | Dec 2022 – Dec 2024',
                   'Performed daily cash and trade reconciliations.',
                   'Second Corporation | Bengaluru | Nov 2021 – Nov 2022',
                   'Data Researcher I | Nov 2021 – Nov 2022', 'Analyzed bond structures.',
                   'Education', 'MBA Finance 2020 – 2022', 'Example Institute',
                   'Certifications', 'Financial Analytics Certificate'])
    response = engine.detect(req)
    sections = {s.category:s for s in response.sections}
    assert sections['summary'].source_heading == 'Profile Snapshot'
    assert sections['skills'].source_heading == 'Core Competencies'
    experience = sections['experience']
    assert len(experience.readable_blocks) == 2
    first, second = experience.readable_blocks
    assert 'Example India' in first.title
    assert len(first.children) == 2
    assert 'Senior Analyst' in first.children[0].title
    assert 'Analyst' in first.children[1].title
    assert len(second.children) == 1
    responsibilities = ' '.join(experience.blocks[i].text for i in first.children[0].block_indices)
    assert 'Managed reconciliations' in responsibilities and 'resolving discrepancies' in responsibilities
    assert 'Built ETL' in responsibilities
    assert 'Python SQL Python' in ''.join(b.text for b in sections['skills'].blocks)
    assert_preserved(req, response)
    def refs(group):
        return group.block_indices + [i for child in group.children for i in refs(child)]
    for section in response.sections:
        actual = [i for group in section.readable_blocks for i in refs(group)]
        headings = [i for i,b in enumerate(section.blocks) if b.assignment in {'section_heading','structural_heading'}]
        expected = [i for i in range(len(section.blocks)) if i not in headings[:1]]
        assert sorted(actual) == expected
        assert len(actual) == len(set(actual))


def test_wrapped_employer_does_not_attach_to_previous_job(engine):
    req = request(['Work History', 'Alpha Private Limited | Jan 2020 – Jan 2022',
                   'Analyst | Jan 2020 – Jan 2022', 'Built reporting systems.',
                   'International Exchange (IX) /', 'NYSE | London | Feb 2022 – Present',
                   'Senior Analyst | Feb 2022 – Present', 'Managed daily reporting.'])
    response = engine.detect(req)
    section = next(s for s in response.sections if s.category == 'experience')
    first, second = section.readable_blocks
    assert 'International Exchange' in second.title and 'NYSE' in second.title
    assert len(second.children) == 1
    def refs(g):
        return g.block_indices + [i for child in g.children for i in refs(child)]
    assert not any('International Exchange' in section.blocks[i].text for i in refs(first))
    assert any('International Exchange' in section.blocks[i].text for i in refs(second))
    assert_preserved(req, response)


def test_preamble_name_and_headline_stay_with_explicit_contact(engine):
    req = request(['Morgan Taylor', 'Senior Analyst — Fund Accounting and Data Operations',
                   'morgan@example.com | London', 'Profile Snapshot', 'Results-driven finance analyst supporting global', 'funds across the trade lifecycle.'],
                  types=['Title', 'NarrativeText', 'NarrativeText', 'Title', 'NarrativeText', 'NarrativeText'])
    result = engine.detect(req)
    personal = next(s for s in result.sections if s.category == 'personal_information')
    assert [b.text for b in personal.blocks] == [e.text for e in req.elements[:3]]
    assert not result.needs_review
    summary = next(s for s in result.sections if s.category == 'summary')
    assert req.elements[4].text in [b.text for b in summary.blocks]
    assert_preserved(req, result)


def test_pdf_columns_merged_headings_and_full_width_continuation(engine):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    from local_api.services.text_extraction_service import extract_text
    class Canvas:
        def __init__(self, stream, pagesize):
            self.stream, self.size, self.commands = stream, pagesize, []
            self.writer = PdfWriter()
        def setFont(self, font, size):
            self.font, self.font_size = ('F2' if 'Bold' in font else 'F1'), size
        def drawString(self, x, y, text):
            escaped = text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
            self.commands.append(f'BT /{self.font} {self.font_size} Tf {x} {y} Td ({escaped}) Tj ET')
        def showPage(self):
            page = self.writer.add_blank_page(*self.size)
            fonts = DictionaryObject()
            for key, face in [('F1', 'Helvetica'), ('F2', 'Helvetica-Bold')]:
                font = DictionaryObject({NameObject('/Type'):NameObject('/Font'), NameObject('/Subtype'):NameObject('/Type1'), NameObject('/BaseFont'):NameObject('/'+face)})
                fonts[NameObject('/'+key)] = self.writer._add_object(font)
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):fonts})
            contents = DecodedStreamObject(); contents.set_data('\n'.join(self.commands).encode('ascii'))
            page[NameObject('/Contents')] = self.writer._add_object(contents)
            self.commands = []
        def save(self):
            self.showPage(); self.writer.write(self.stream)
    stream = BytesIO()
    canvas = Canvas(stream, pagesize=(600, 840))
    def line(x, y, text, heading=False):
        canvas.setFont('Helvetica-Bold' if heading else 'Helvetica', 16 if heading else 9)
        canvas.drawString(x, y, text)
    line(25, 790, 'Morgan Taylor', True)
    line(25, 740, 'Contact', True); line(25, 722, 'morgan@example.com')
    line(230, 740, 'Profile Snapshot', True); line(230, 722, 'Experienced banking transformation specialist.')
    line(25, 650, 'Skills', True); line(25, 632, 'Project management'); line(25, 618, 'SQL and communication')
    line(230, 650, 'Core Industry & Domain Expertise', True); line(230, 632, 'Retail lending and regulatory compliance')
    line(230, 560, 'Business Analysis & Product Management', True); line(230, 542, 'Requirement elicitation and backlog ownership')
    line(25, 280, 'Certifications', True); line(25, 262, 'Agile coaching certificate')
    canvas.showPage()
    line(25, 790, 'Program management certificate')
    line(230, 740, 'Work History', True)
    line(230, 714, 'Novus Bank Ltd | Jan 2020 - Dec 2022')
    line(230, 699, 'Senior Analyst'); line(230, 678, 'Built credit reconciliation services.')
    line(25, 110, 'Next Systems Ltd | Jan 2023 - Present')
    line(25, 94, 'Lead Analyst')
    canvas.showPage()
    line(25, 790, 'Led enterprise-wide delivery across lending, payments and regulatory reporting.')
    line(25, 610, 'Tools & Technology', True); line(25, 590, 'Professional'); line(25, 574, 'Python SQL Docker')
    line(25, 500, 'Education', True); line(25, 480, 'MBA Finance, Example University, 2018')
    line(25, 40, 'Place: London')
    canvas.save()
    extracted = extract_text(stream.getvalue(), 'synthetic-columns.pdf', strategy='fast')
    assert any(e.metadata.get('pdf_source_lines') for e in extracted.elements)
    req = SectionRequest(elements=extracted.elements, extraction_verified=True)
    response = engine.detect(req)
    sections = {s.category:s for s in response.sections}
    expected = {'personal_information': ['Morgan Taylor', 'Place: London'],
                'summary': ['Experienced banking transformation specialist.'],
                'skills': ['Project management', 'Retail lending', 'Requirement elicitation', 'Python SQL Docker'],
                'certifications': ['Agile coaching certificate', 'Program management certificate'],
                'experience': ['Novus Bank', 'Next Systems', 'Led enterprise-wide delivery'],
                'education': ['MBA Finance']}
    for category, markers in expected.items():
        for marker in markers:
            assert marker in sections[category].text, (category, marker)
            assert not any(marker in s.text for s in response.sections if s.category != category)
    assert 'projects' not in sections
    assert 'Program management certificate' not in sections['experience'].text
    assert_preserved(req, response)


def test_canonical_text_merges_wordpiece_spans_without_inserting_spaces(engine):
    req = request(['Work History\n' + 'Extraordinary' * 240 + '\nEducation\nMBA Finance'])
    response = engine.detect(req)
    experience = next(s for s in response.sections if s.category == 'experience')
    assert 'Extraordinary' * 240 in experience.text
    assert 'MBA Finance' not in experience.text
    assert len({s.category for s in response.sections}) == len(response.sections)
    assert_preserved(req, response)
